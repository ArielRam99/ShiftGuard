from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


MIME_TYPES = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _append_excel_row(worksheet, values):
    """Prevent user-provided spreadsheet values from becoming formulas."""
    safe_values = [
        f"'{value}"
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@"))
        else value
        for value in values
    ]
    worksheet.append(safe_values)


def _style_worksheet(worksheet):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    worksheet.freeze_panes = "A2"
    for column in worksheet.columns:
        letter_name = column[0].column_letter
        width = min(50, max(12, *(len(str(cell.value or "")) + 2 for cell in column)))
        worksheet.column_dimensions[letter_name].width = width


def _pdf_table(rows, widths=None):
    table = Table(rows, repeatRows=1, colWidths=widths)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EAF2F8")]),
            ]
        )
    )
    return table


def schedule_report(shifts, file_format):
    if file_format == "xlsx":
        return _schedule_xlsx(shifts)
    return _schedule_pdf(shifts)


def _schedule_xlsx(shifts):
    workbook = Workbook()
    schedule_sheet = workbook.active
    schedule_sheet.title = "Schedules"
    _append_excel_row(
        schedule_sheet,
        [
            "Shift ID",
            "Date",
            "Start",
            "End",
            "Role",
            "Recommended Staff",
            "Range",
            "Model",
            "Status",
            "Manager",
        ]
    )
    for shift in shifts:
        staffing_range = shift["staffing_range"]
        _append_excel_row(
            schedule_sheet,
            [
                shift["id"],
                shift["shift_date"],
                shift["start_time"],
                shift["end_time"],
                shift["required_role"],
                shift["required_staff"],
                f"{staffing_range['minimum']}-{staffing_range['maximum']}",
                shift["model_source"],
                shift["status"],
                shift["decided_by"] or "",
            ]
        )
    _style_worksheet(schedule_sheet)

    assignment_sheet = workbook.create_sheet("Assignments")
    _append_excel_row(
        assignment_sheet,
        [
            "Shift ID",
            "Employee ID",
            "Employee",
            "Role",
            "Projected Hours",
            "Overtime Hours",
            "Status",
            "Recommendation Reason",
        ]
    )
    for shift in shifts:
        for assignment in shift["assignments"]:
            _append_excel_row(
                assignment_sheet,
                [
                    shift["id"],
                    assignment["employee_id"],
                    assignment["employee_name"],
                    assignment["role"],
                    assignment["projected_weekly_hours"],
                    assignment["projected_overtime_hours"],
                    assignment["status"],
                    assignment["reason"],
                ]
            )
    _style_worksheet(assignment_sheet)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _schedule_pdf(shifts):
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        leftMargin=0.35 * inch,
        rightMargin=0.35 * inch,
        topMargin=0.4 * inch,
        bottomMargin=0.4 * inch,
        title="ShiftGuard AI Schedule Report",
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("ShiftGuard AI Schedule Report", styles["Title"]),
        Paragraph(
            "AI recommendations are advisory and require manager approval.",
            styles["BodyText"],
        ),
        Spacer(1, 0.15 * inch),
    ]
    rows = [["ID", "Date", "Time", "Role", "Staff", "Range", "Model", "Status", "Employees"]]
    for shift in shifts:
        employees = ", ".join(item["employee_name"] for item in shift["assignments"]) or "Unfilled"
        staffing_range = shift["staffing_range"]
        rows.append(
            [
                shift["id"],
                shift["shift_date"],
                f"{shift['start_time']}-{shift['end_time']}",
                shift["required_role"],
                shift["required_staff"],
                f"{staffing_range['minimum']}-{staffing_range['maximum']}",
                shift["model_source"],
                shift["status"],
                employees,
            ]
        )
    story.append(
        _pdf_table(
            rows,
            [
                0.35 * inch,
                0.75 * inch,
                0.85 * inch,
                0.8 * inch,
                0.4 * inch,
                0.45 * inch,
                0.9 * inch,
                0.65 * inch,
                2.8 * inch,
            ],
        )
    )
    document.build(story)
    output.seek(0)
    return output


def analytics_report(summary, comparison, file_format):
    if file_format == "xlsx":
        return _analytics_xlsx(summary, comparison)
    return _analytics_pdf(summary, comparison)


def _analytics_xlsx(summary, comparison):
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    _append_excel_row(summary_sheet, ["Metric", "Value"])
    for key, value in summary.items():
        _append_excel_row(
            summary_sheet, [key.replace("_", " ").title(), value]
        )
    _style_worksheet(summary_sheet)

    model_sheet = workbook.create_sheet("Model Comparison")
    _append_excel_row(
        model_sheet,
        ["Strategy", "MAE", "RMSE", "R2", "Confidence Error Margin", "Best Model"],
    )
    for result in comparison["models"]:
        _append_excel_row(
            model_sheet,
            [
                result["strategy"],
                result["mae"],
                result["rmse"],
                result["r2"],
                result["confidence_error_margin"],
                result["strategy"] == comparison["best_model"],
            ]
        )
    _style_worksheet(model_sheet)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _analytics_pdf(summary, comparison):
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        title="ShiftGuard AI Analytics Report",
    )
    styles = getSampleStyleSheet()
    story = [Paragraph("ShiftGuard AI Analytics Report", styles["Title"])]
    summary_rows = [["Metric", "Value"]] + [
        [key.replace("_", " ").title(), value] for key, value in summary.items()
    ]
    story.extend([_pdf_table(summary_rows), Spacer(1, 0.25 * inch)])
    model_rows = [["Strategy", "MAE", "RMSE", "R2", "Range Margin", "Best"]]
    for result in comparison["models"]:
        model_rows.append(
            [
                result["strategy"],
                result["mae"],
                result["rmse"],
                result["r2"],
                result["confidence_error_margin"],
                "Yes" if result["strategy"] == comparison["best_model"] else "",
            ]
        )
    if comparison["models"]:
        story.append(_pdf_table(model_rows))
    else:
        story.append(
            Paragraph(
                "Not enough historical records to compare machine-learning models.",
                styles["BodyText"],
            )
        )
    document.build(story)
    output.seek(0)
    return output
